"""把筛选报告 Markdown 导出为 docx（标题/段落/项目符号/表格/图片）。

只处理本模块自己生成的受限 Markdown 子集，不做通用解析。
"""

from pathlib import Path

import docx
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


def _set_cjk_font(doc) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _add_runs(p, text: str) -> None:
    """支持 **加粗** 的行内标记。"""
    for i, seg in enumerate(text.split("**")):
        run = p.add_run(seg)
        run.bold = (i % 2 == 1)


def _add_table(doc, block: list[str]) -> None:
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in block]
    if len(rows) < 2:
        return
    header, body = rows[0], rows[2:] if set(rows[1][0]) <= {"-", ":"} else rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        _add_runs(cell.paragraphs[0], c)
        for r in cell.paragraphs[0].runs:
            r.bold = True
    for row in body:
        cells = t.add_row().cells
        for i, c in enumerate(row[:len(header)]):
            cells[i].text = ""
            _add_runs(cells[i].paragraphs[0], c)


def convert(md_path: Path, out_path: Path, images: list[Path]) -> Path:
    doc = docx.Document()
    _set_cjk_font(doc)
    by_stem = {p.stem: p for p in images}
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _add_table(doc, block)
            continue
        if s.startswith("### "):
            title = s[4:]
            doc.add_heading(title, level=3)
            i += 1
            if title.startswith("图表"):
                # 每个“### 图表”节按其中列出的 "- `xxx.png`：说明" 顺序插图
                while i < len(lines) and not lines[i].strip().startswith("#"):
                    item = lines[i].strip()
                    if item.startswith("- `"):
                        name, _, desc = item[3:].partition("`")
                        img = by_stem.get(Path(name).stem)
                        if img is not None and img.exists():
                            doc.add_picture(str(img), width=Inches(6.0))
                            doc.add_paragraph(desc.lstrip("：: ") or img.stem)
                    i += 1
            continue
        if s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith("> "):
            p = doc.add_paragraph()
            _add_runs(p, s[2:])
            p.italic = True
        elif s.startswith("- "):
            _add_runs(doc.add_paragraph(style="List Bullet"), s[2:])
        elif s in ("", "---"):
            pass
        else:
            _add_runs(doc.add_paragraph(), s)
        i += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
